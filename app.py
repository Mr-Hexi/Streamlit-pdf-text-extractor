import io
import os
import re
import shutil
from typing import Any, Dict, List

import openpyxl
import pdfplumber
import pytesseract
import streamlit as st
from pdf2image import convert_from_bytes
from pdf2image.exceptions import PDFInfoNotInstalledError, PDFPageCountError
from pdfminer.pdfdocument import PDFPasswordIncorrect
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image

try:
    import fitz
except ImportError:
    fitz = None

# ─── Page config ──────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Bank Statement → Excel",
    page_icon="🏦",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ─── Styles ───────────────────────────────────────────────────────────────────
st.markdown("""
<style>
    /* Hide default Streamlit chrome */
    #MainMenu, footer { visibility: hidden; }
    .block-container { padding-top: 2rem; padding-bottom: 2rem; }

    /* Header */
    .app-header {
        background: #0F1F35;
        border-bottom: 3px solid #63B3ED;
        padding: 1.1rem 1.5rem;
        border-radius: 10px;
        margin-bottom: 1.5rem;
        display: flex;
        align-items: center;
        gap: 1rem;
    }
    .app-header h1 {
        color: #FFFFFF;
        font-size: 1.4rem;
        margin: 0;
        font-weight: 700;
    }
    .app-header p {
        color: #94A3B8;
        font-size: 0.8rem;
        margin: 0;
        font-family: monospace;
    }

    /* Stat cards */
    .stat-row { display: flex; gap: 1rem; margin-bottom: 1.5rem; flex-wrap: wrap; }
    .stat-card {
        background: #FFFFFF;
        border: 1px solid #CBD5E1;
        border-radius: 10px;
        padding: 1rem 1.5rem;
        flex: 1;
        min-width: 160px;
        box-shadow: 0 1px 3px rgba(0,0,0,0.06);
    }
    .stat-label { font-size: 0.75rem; color: #64748B; font-weight: 600; text-transform: uppercase; letter-spacing: 0.04em; }
    .stat-value { font-size: 1.5rem; font-weight: 700; color: #0F1F35; margin-top: 0.2rem; }
    .stat-sub   { font-size: 0.75rem; color: #94A3B8; font-family: monospace; }

    /* Info grid */
    .info-grid { display: grid; grid-template-columns: 1fr 1fr; gap: 0; border: 1px solid #CBD5E1; border-radius: 8px; overflow: hidden; }
    .info-row  { display: contents; }
    .info-label-cell { background: #F8FAFC; padding: 0.55rem 0.9rem; font-size: 0.85rem; font-weight: 600; color: #0F1F35; border-bottom: 1px solid #E2E8F0; }
    .info-value-cell { background: #FFFFFF;  padding: 0.55rem 0.9rem; font-size: 0.85rem; color: #334155; border-bottom: 1px solid #E2E8F0; font-family: monospace; }

    /* Prompt box */
    .prompt-box {
        background: #0F1F35;
        border-radius: 10px;
        padding: 1.25rem;
        color: #94A3B8;
        font-family: monospace;
        font-size: 0.78rem;
        white-space: pre-wrap;
        word-break: break-word;
        max-height: 300px;
        overflow-y: auto;
        border: 1px solid #2B5282;
    }

    /* Success / warning banners */
    .success-pill {
        display: inline-flex; align-items: center; gap: 0.4rem;
        background: #F0FFF4; border: 1px solid #9AE6B4;
        color: #276749; border-radius: 20px; padding: 0.3rem 0.9rem;
        font-size: 0.8rem; font-weight: 600; margin-bottom: 1rem;
    }
    .warn-pill {
        display: inline-flex; align-items: center; gap: 0.4rem;
        background: #FFFBEB; border: 1px solid #F6E05E;
        color: #744210; border-radius: 20px; padding: 0.3rem 0.9rem;
        font-size: 0.8rem; font-weight: 600; margin-bottom: 1rem;
    }
</style>
""", unsafe_allow_html=True)

# ─── AI Extraction Prompt ─────────────────────────────────────────────────────
AI_EXCEL_PROMPT_TEMPLATE = """You are an expert bank statement data extraction assistant.

Create a clean Excel-ready output from the PDF text below.

Requirements:
- Preserve every transaction row in the statement.
- Create two sheets: Account Details and Transactions.
- Account Details columns: Field, Value.
- Transactions columns: Date, Particulars, Withdrawal, Deposit, Balance.
- Keep amounts exactly as written unless a clear formatting cleanup is needed.
- Put withdrawals and deposits in separate columns. Leave missing values blank.
- If a transaction description wraps across lines, merge it into the same row.
- Ignore page headers, footers, disclaimers, duplicate column headings, and summary noise.
- Return the result as CSV blocks, one block per sheet, so I can paste/import it into Excel.
- If anything is uncertain, add a short Notes section after the CSV blocks.

PDF text:
{pdf_text}
"""

DATE_RE   = re.compile(r'^(\d{2}-(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)-\d{4})', re.IGNORECASE)
NUMBER_RE = re.compile(r'^[\d,]+(?:\.\d+)?$')

DISCLAIMER_MARKERS = [
    'this is an auto generated e-statement',
    'this is auto generated',
    'for any query',
    'if you need any clarification',
    'please contact your branch',
    'please contact the branch',
    'generated by',
    'statement generated',
    'disclaimer',
]

# ─── Core Helper Functions ────────────────────────────────────────────────────
class PdfPasswordError(Exception):
    pass


class CanaraOcrSetupError(Exception):
    pass


def normalize_pdf_password(password: str | None) -> str | None:
    if password is None:
        return None
    password = password.strip()
    return password or None


def safe_float(val: str) -> float:
    try:
        return float(val.replace(",", ""))
    except (ValueError, AttributeError):
        return 0.0


def open_pdf(file_bytes: bytes, password: str | None = None):
    password = normalize_pdf_password(password)
    try:
        return pdfplumber.open(io.BytesIO(file_bytes), password=password)
    except PDFPasswordIncorrect as exc:
        if password:
            message = "The PDF password is incorrect. Please check it and try again."
        else:
            message = "This PDF is password protected. Enter the PDF password and try again."
        raise PdfPasswordError(message) from exc


def find_poppler_path() -> str | None:
    if shutil.which("pdfinfo") and shutil.which("pdftoppm"):
        return None

    env_path = os.environ.get("POPPLER_PATH")
    candidates = [
        env_path,
        os.path.join(os.getcwd(), "poppler", "bin"),
        os.path.join(os.getcwd(), "poppler", "Library", "bin"),
        r"C:\poppler\bin",
        r"C:\poppler\Library\bin",
        r"C:\Program Files\poppler\bin",
        r"C:\Program Files\poppler\Library\bin",
    ]
    for path in candidates:
        if not path:
            continue
        if os.path.exists(os.path.join(path, "pdfinfo.exe")) and os.path.exists(os.path.join(path, "pdftoppm.exe")):
            return path

    return None


def get_poppler_help_message() -> str:
    return (
        "Canara Bank extraction needs a PDF rendering backend. Install PyMuPDF with "
        "`pip install PyMuPDF`, or install Poppler and add its bin folder to PATH/set POPPLER_PATH. "
        "For Poppler, the folder must contain pdfinfo.exe and pdftoppm.exe."
    )


def find_tesseract_path() -> str | None:
    if shutil.which("tesseract"):
        return None
    candidates = [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        os.path.join(os.environ.get("LOCALAPPDATA", ""), r"Programs\Tesseract-OCR\tesseract.exe"),
        os.path.join(os.environ.get("LOCALAPPDATA", ""), r"Tesseract-OCR\tesseract.exe"),
        os.path.join(os.environ.get("USERPROFILE", ""), r"AppData\Local\Programs\Tesseract-OCR\tesseract.exe"),
        os.path.join(os.environ.get("USERPROFILE", ""), r"AppData\Local\Tesseract-OCR\tesseract.exe"),
    ]
    for path in candidates:
        if path and os.path.exists(path):
            return path
    return None


# Configure tesseract path on startup if detected in default folders
_tesseract_path = find_tesseract_path()
if _tesseract_path:
    pytesseract.pytesseract.tesseract_cmd = _tesseract_path


def get_tesseract_help_message() -> str:
    return (
        "Canara Bank extraction requires Tesseract OCR to read text from image-based PDFs.\n\n"
        "**Please install Tesseract OCR to resolve this:**\n\n"
        "1. Download the Windows installer from the UB Mannheim community build repository:\n"
        "   https://github.com/UB-Mannheim/tesseract/wiki\n"
        "2. Run the installer (e.g., `tesseract-ocr-w64-setup-v5.x.x.exe`) and complete the installation "
        "using the default location (`C:\\Program Files\\Tesseract-OCR`).\n"
        "3. Once installed, restart this Streamlit application. The app will automatically detect Tesseract at the default installation path, or you can add it to your system PATH."
    )


def is_disclaimer_text(text: str) -> bool:
    if not text:
        return False
    normalized = re.sub(r'\s+', ' ', text).strip().lower()
    return any(m in normalized for m in DISCLAIMER_MARKERS)

def starts_new_transaction_row(line_words: list[dict]) -> bool:
    if not line_words:
        return False
    return bool(DATE_RE.match(line_words[0].get('text', '').strip()))

def words_to_lines(page, y_tolerance: int = 6) -> list[list[dict]]:
    words = page.extract_words(x_tolerance=4, y_tolerance=4)
    if not words:
        return []
    words.sort(key=lambda w: (w['top'] + w['bottom']) / 2)
    lines, current_line, current_y = [], [], None
    for w in words:
        word_mid = (w['top'] + w['bottom']) / 2
        if current_y is None:
            current_y, current_line = word_mid, [w]
        elif starts_new_transaction_row([w]) and current_line:
            lines.append(sorted(current_line, key=lambda x: x['x0']))
            current_line, current_y = [w], word_mid
        elif abs(word_mid - current_y) <= y_tolerance:
            current_line.append(w)
            current_y = sum((cw['top'] + cw['bottom']) / 2 for cw in current_line) / len(current_line)
        else:
            lines.append(sorted(current_line, key=lambda x: x['x0']))
            current_line, current_y = [w], word_mid
    if current_line:
        lines.append(sorted(current_line, key=lambda x: x['x0']))
    return lines

def parse_transaction_line(line_words: list[dict]) -> dict | None:
    if not line_words:
        return None
    first_word = line_words[0]['text']
    if not DATE_RE.match(first_word):
        return None

    date = first_word
    particulars_words, numbers = [], []

    for w in line_words[1:]:
        is_numeric = NUMBER_RE.match(w['text'].replace(',', ''))
        if w['x1'] > 360 and is_numeric:
            numbers.append((w['x1'], w['text']))
        else:
            particulars_words.append(w['text'])

    particulars = " ".join(particulars_words).strip()
    withdrawal, deposit, balance = "", "", ""

    if len(numbers) >= 1:
        balance = numbers[-1][1]
    if len(numbers) >= 2:
        x1, val = numbers[-2]
        if x1 < 440:
            withdrawal = val
        else:
            deposit = val
    if len(numbers) >= 3:
        x1, val = numbers[-3]
        if x1 < 440:
            withdrawal = val
        else:
            deposit = val

    return {"date": date, "particulars": particulars,
            "withdrawal": withdrawal, "deposit": deposit, "balance": balance}

# ─── Bank Specific Parsers ────────────────────────────────────────────────────
def extract_uco_account_info(first_page_text: str) -> dict:
    info = {}
    patterns = {
        "Statement Period": r'Between\s+(\d{2}-\d{2}-\d{4}\s+and\s+\d{2}-\d{2}-\d{4})',
        "Account Number":   r'account number\s+(\d+)',
        "Customer ID":      r'Customer ID\s+([A-Z0-9]+)',
        "CKYC ID":          r'CKYC ID\s+(\d+)',
        "Account Type":     r'A/c Type\s+(\w+)',
        "Mobile No":        r'Mobile No[.\s]+(\d+)',
        "Email (Customer)": r'E-Mail ID\s+([A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,})',
        "Branch Code":      r'Branch Code\s+(\d+)',
        "Branch Name":      r'Branch Name\s+([A-Z]+)',
        "IFSC Code":        r'IFSC Code\s+([A-Z0-9]+)',
        "MICR Code":        r'MICR Code\s+(\d+)',
    }
    for field, pattern in patterns.items():
        m = re.search(pattern, first_page_text, re.IGNORECASE)
        info[field] = m.group(1).strip() if m else ""

    name_m = re.search(r'((?:[A-Z][A-Z\s&\./]+\n)+)Name\n', first_page_text)
    if name_m:
        lines = [l.strip() for l in name_m.group(1).strip().split('\n') if l.strip()]
        lines = [l for l in lines if l != info.get("Branch Name", "")]
        info["Account Name"] = " / ".join(lines)

    addr_blocks = list(re.finditer(r'Address\n', first_page_text))
    if len(addr_blocks) >= 2:
        start = addr_blocks[-1].end()
        end_m = re.search(r'IFSC Code|MICR Code|A/c Type', first_page_text[start:])
        addr_raw = first_page_text[start:start + end_m.start()] if end_m else first_page_text[start:start + 200]
        addr_lines = [l.strip() for l in addr_raw.split('\n') if l.strip()]
        info["Address"] = ", ".join(addr_lines)
    elif len(addr_blocks) == 1:
        start = addr_blocks[0].end()
        end_m = re.search(r'IFSC Code|MICR Code|A/c Type', first_page_text[start:])
        addr_raw = first_page_text[start:start + end_m.start()] if end_m else first_page_text[start:start + 200]
        addr_lines = [l.strip() for l in addr_raw.split('\n') if l.strip()]
        info["Address"] = ", ".join(addr_lines)

    return info

def extract_uco_transactions(file_bytes: bytes, password: str | None = None) -> list[dict]:
    transactions = []
    with open_pdf(file_bytes, password) as pdf:
        for page in pdf.pages:
            lines = words_to_lines(page)
            for line_words in lines:
                line_text = " ".join(w['text'] for w in line_words)
                if is_disclaimer_text(line_text):
                    continue
                if 'Opening Balance' in line_text:
                    continue
                tx = parse_transaction_line(line_words)
                if tx:
                    transactions.append(tx)
                elif transactions and line_words and not starts_new_transaction_row(line_words):
                    continuation = [w['text'] for w in line_words if w['x0'] < 360 and not DATE_RE.match(w['text'])]
                    if continuation:
                        extra = " ".join(continuation).strip()
                        if extra and not is_disclaimer_text(extra):
                            if transactions[-1]['particulars']:
                                transactions[-1]['particulars'] += " " + extra
                            else:
                                transactions[-1]['particulars'] = extra
    return transactions

# ─── Router Logic ─────────────────────────────────────────────────────────────
# Canara Bank Parser
_CB_X_DATE_MAX        = 340
_CB_X_PARTICULARS_MAX = 880
_CB_X_DEPOSITS_MAX    = 1140
_CB_X_WITHDRAWALS_MAX = 1440

_CB_DATE_RE   = re.compile(r'^\d{2}[-/.\s]\d{2}[-/.\s]\d{2,4}$')
_CB_NUMBER_RE = re.compile(r'^\d[\d,]*(?:\.\d{1,2})?$')
_CB_CHQ_RE    = re.compile(r'^chq[;:.]?$', re.IGNORECASE)
_CB_NOISE_RE  = re.compile(
    r'^(date|particulars|deposits?|withdrawals?|balance|opening|closing|'
    r'constituent|ombudsman|phish|unauthori|canara|branch|account|'
    r'rbi|atm|pin|intents|deemed|correct|requested|always|login|'
    r'click|note|please|beware|attempt|holder|along)$',
    re.IGNORECASE
)


def _cb_pdf_images(file_bytes: bytes, password: str | None = None, **kwargs):
    password = normalize_pdf_password(password)
    pymupdf_images = _cb_pdf_images_with_pymupdf(file_bytes, password, **kwargs)
    if pymupdf_images is not None:
        return pymupdf_images

    poppler_path = find_poppler_path()
    try:
        return convert_from_bytes(
            file_bytes,
            dpi=200,
            userpw=password,
            poppler_path=poppler_path,
            **kwargs,
        )
    except PDFInfoNotInstalledError as exc:
        raise CanaraOcrSetupError(get_poppler_help_message()) from exc
    except PDFPageCountError as exc:
        message = str(exc)
        if "Unable to get page count" in message:
            raise CanaraOcrSetupError(get_poppler_help_message()) from exc
        raise


def _cb_pdf_images_with_pymupdf(file_bytes: bytes, password: str | None = None, **kwargs) -> list[Image.Image] | None:
    if fitz is None:
        return None

    first_page = kwargs.get("first_page")
    last_page = kwargs.get("last_page")
    dpi = kwargs.get("dpi", 200)
    scale = dpi / 72
    matrix = fitz.Matrix(scale, scale)

    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        if doc.needs_pass:
            if not password or not doc.authenticate(password):
                raise PdfPasswordError("The PDF password is incorrect. Please check it and try again.")

        start = max((first_page or 1) - 1, 0)
        end = min(last_page or doc.page_count, doc.page_count)
        images = []

        for page_index in range(start, end):
            page = doc.load_page(page_index)
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            images.append(img)

        doc.close()
        return images
    except PdfPasswordError:
        raise
    except Exception:
        return None


def _cb_confidence_ok(confidence: str) -> bool:
    try:
        return float(confidence) > 20
    except ValueError:
        return False


def _cb_normalize_date(text: str) -> str:
    normalized = text.strip().replace("O", "0").replace("o", "0")
    normalized = re.sub(r'[-/.\s]+', '-', normalized)
    if not _CB_DATE_RE.match(text.strip()) and not re.match(r'^\d{2}-\d{2}-\d{2,4}$', normalized):
        return ""
    day, month, year = normalized.split("-")
    if len(year) == 2:
        year = "20" + year
    return f"{day}-{month}-{year}"


def _cb_normalize_amount(text: str) -> str:
    normalized = re.sub(r'[^\d,.]', '', text.strip())
    normalized = normalized.strip(',.')
    if not _CB_NUMBER_RE.match(normalized):
        return ""
    return normalized


def _cb_parse_page(img) -> list[dict]:
    data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
    words = [
        {'text': t.strip(), 'x': data['left'][i], 'y': data['top'][i]}
        for i, t in enumerate(data['text'])
        if t.strip() and _cb_confidence_ok(data['conf'][i])
    ]
    if not words:
        return []

    dates = []
    for w in words:
        date = _cb_normalize_date(w['text'])
        if date and w['x'] < _CB_X_DATE_MAX:
            dates.append((w['y'], date))
    if not dates:
        return []

    chq_ys = sorted(set(w['y'] for w in words if _CB_CHQ_RE.match(w['text'])))
    img_height = img.size[1]
    transactions = []

    for idx, (date_y, date_str) in enumerate(dates):
        own_chq_y = next((y for y in chq_ys if y >= date_y), img_height)
        if idx == 0:
            header_ys = [w['y'] for w in words if w['text'].lower() == 'particulars']
            part_start_y = (header_ys[0] + 20) if header_ys else 0
        else:
            prev_chq_ys = [y for y in chq_ys if y >= dates[idx - 1][0]]
            part_start_y = (prev_chq_ys[0] + 30) if prev_chq_ys else dates[idx - 1][0]

        next_date_y = dates[idx + 1][0] if idx + 1 < len(dates) else img_height
        particulars_parts, deposits, withdrawals, balances = [], [], [], []

        for w in words:
            x, y, t = w['x'], w['y'], w['text']
            if x < _CB_X_DATE_MAX:
                continue
            elif x < _CB_X_PARTICULARS_MAX:
                if part_start_y <= y < own_chq_y and not _CB_NOISE_RE.match(t) and len(t) > 1:
                    particulars_parts.append((y, x, t))
            elif date_y <= y < next_date_y:
                amount = _cb_normalize_amount(t)
                if x < _CB_X_DEPOSITS_MAX:
                    if amount:
                        deposits.append(amount)
                elif x < _CB_X_WITHDRAWALS_MAX:
                    if amount:
                        withdrawals.append(amount)
                else:
                    if amount:
                        balances.append(amount)

        if not balances:
            continue

        particulars_parts.sort(key=lambda w: (w[0], w[1]))
        transactions.append({
            'date':        date_str,
            'particulars': " ".join(t for _, _, t in particulars_parts),
            'deposit':     deposits[-1]    if deposits    else "",
            'withdrawal':  withdrawals[-1] if withdrawals else "",
            'balance':     balances[-1],
        })

    return transactions


def extract_canara_account_info(ocr_text: str) -> dict:
    info = {}
    patterns = {
        "Statement Period": r'between\s+(\d{2}-\w{3}-\d{4}\s+and\s+\d{2}-\w{3}-\d{4})',
        "Account Number":   r'A/c\s+(X+\d+)',
        "Customer ID":      r'Customer Id\s+(\S+)',
        "Account Name":     r'Name\s+([A-Z][A-Z\s]+?)(?:\n\nPhone|\nPhone)',
        "Phone":            r'Phone\s+(\+?\d+)',
        "Branch Code":      r'Branch Code\s*\nBranch Name\s*\nIFSC Code\s*\n\nAddress\s*\nRoad\s*\n\n(\d+)',
        "Branch Name":      r'Branch Code\s*\nBranch Name\s*\nIFSC Code\s*\n\nAddress\s*\nRoad\s*\n\n\d+\s*\n\n([A-Z]+)',
        "IFSC Code":        r'(CNRB\w+|[A-Z]{4}0\w+)',
    }
    for field, pattern in patterns.items():
        m = re.search(pattern, ocr_text, re.IGNORECASE)
        info[field] = m.group(1).strip() if m else ""

    addr_m = re.search(r'Address\s+(.*?)(?=\n\nBranch Code)', ocr_text, re.DOTALL)
    if addr_m:
        lines = [l.strip() for l in addr_m.group(1).split('\n') if l.strip()]
        info["Address"] = ", ".join(lines)

    return info


def extract_canara_transactions(file_bytes: bytes, password: str | None = None) -> list[dict]:
    images = _cb_pdf_images(file_bytes, password)
    try:
        return [tx for img in images for tx in _cb_parse_page(img)]
    except (pytesseract.TesseractNotFoundError, pytesseract.TesseractError) as exc:
        raise CanaraOcrSetupError(get_tesseract_help_message()) from exc


def extract_canara_info_from_bytes(file_bytes: bytes, password: str | None = None) -> dict:
    images = _cb_pdf_images(file_bytes, password, first_page=1, last_page=1)
    if not images:
        raise CanaraOcrSetupError("Could not render the first page of the PDF.")
    try:
        ocr_text = pytesseract.image_to_string(images[0])
    except (pytesseract.TesseractNotFoundError, pytesseract.TesseractError) as exc:
        raise CanaraOcrSetupError(get_tesseract_help_message()) from exc
    return extract_canara_account_info(ocr_text)


def parse_bank_statement(bank_name: str, file_bytes: bytes, page1_text: str, password: str | None = None):
    """Routes the PDF to the correct parsing logic based on the selected bank."""
    
    if bank_name == "UCO Bank":
        account_info = extract_uco_account_info(page1_text)
        transactions = extract_uco_transactions(file_bytes, password)
        return account_info, transactions

    elif bank_name == "Canara Bank":
        account_info = extract_canara_info_from_bytes(file_bytes, password)
        transactions = extract_canara_transactions(file_bytes, password)
        return account_info, transactions
        
    elif bank_name == "SBI (Coming Soon)":
        st.warning("SBI parsing logic is not yet implemented. Please check back later!")
        st.stop()
        
    elif bank_name == "HDFC (Coming Soon)":
        st.warning("HDFC parsing logic is not yet implemented. Please check back later!")
        st.stop()
        
    else:
        st.error("Unsupported bank selected.")
        st.stop()


# ─── Standard Utilities ───────────────────────────────────────────────────────
def extract_first_page_text(file_bytes: bytes, password: str | None = None) -> str:
    with open_pdf(file_bytes, password) as pdf:
        return pdf.pages[0].extract_text() or "" if pdf.pages else ""


def extract_pdf_text(file_bytes: bytes, password: str | None = None) -> tuple[str, int]:
    page_texts = []
    with open_pdf(file_bytes, password) as pdf:
        total = len(pdf.pages)
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            page_texts.append(f"--- Page {i} of {total} ---\n{text.strip()}")
    return "\n\n".join(page_texts).strip(), len(page_texts)

def build_excel_workbook(account_info: dict, transactions: list[dict], bank_name: str) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws_info = wb.active
    ws_info.title = "Account Details"

    title_font  = Font(color="1E3A5F", bold=True, name="Calibri", size=14)
    label_font  = Font(bold=True, name="Calibri", size=11)
    value_font  = Font(name="Calibri", size=11)
    thin_border = Border(bottom=Side(style='thin', color='CBD5E1'))

    ws_info.merge_cells("A1:B1")
    c = ws_info["A1"]
    c.value = f"{bank_name} — Account Statement Details"
    c.font  = title_font
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws_info.row_dimensions[1].height = 30
    ws_info.column_dimensions["A"].width = 25
    ws_info.column_dimensions["B"].width = 45

    for r, (key, val) in enumerate(account_info.items(), start=2):
        lc = ws_info.cell(row=r, column=1, value=key)
        vc = ws_info.cell(row=r, column=2, value=val)
        lc.font, vc.font = label_font, value_font
        lc.border, vc.border = thin_border, thin_border
        if r % 2 == 0:
            fill = PatternFill(start_color="EBF8FF", end_color="EBF8FF", fill_type="solid")
            lc.fill, vc.fill = fill, fill

    ws_tx = wb.create_sheet(title="Transactions")
    headers    = ["Date", "Particulars", "Withdrawals (₹)", "Deposits (₹)", "Balance (₹)"]
    col_widths = [15, 55, 18, 18, 18]
    hdr_fill   = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
    hdr_font   = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
    debit_font = Font(name="Calibri", size=10, color="C53030")
    credit_font= Font(name="Calibri", size=10, color="276749")
    normal_font= Font(name="Calibri", size=10)

    for i, (h, w) in enumerate(zip(headers, col_widths), 1):
        c = ws_tx.cell(row=1, column=i, value=h)
        c.fill, c.font = hdr_fill, hdr_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws_tx.column_dimensions[get_column_letter(i)].width = w
    ws_tx.row_dimensions[1].height = 22

    for r, tx in enumerate(transactions, 2):
        row_data = [tx.get("date",""), tx.get("particulars",""),
                    tx.get("withdrawal",""), tx.get("deposit",""), tx.get("balance","")]
        alt = PatternFill(start_color="F0F4F8", end_color="F0F4F8", fill_type="solid") if r % 2 == 0 else None
        for ci, val in enumerate(row_data, 1):
            cell = ws_tx.cell(row=r, column=ci, value=val)
            if alt:
                cell.fill = alt
            if ci == 3 and val:
                cell.font, cell.alignment = debit_font, Alignment(horizontal="right")
            elif ci == 4 and val:
                cell.font, cell.alignment = credit_font, Alignment(horizontal="right")
            elif ci == 5:
                cell.font  = Font(name="Calibri", size=10, bold=True)
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.font = normal_font

    ws_tx.freeze_panes = "A2"
    return wb

# ─── Streamlit UI ─────────────────────────────────────────────────────────────
st.markdown("""
<div class="app-header">
  <div>
    <h1>🏦 Bank Statement → Excel</h1>
    <p>Deterministic extraction · Works offline · Multi-bank support</p>
  </div>
</div>
""", unsafe_allow_html=True)

# ── Mode selector ─────────────────────────────────────────────────────────────
st.markdown("#### What do you want to do?")
mode = st.radio(
    label="mode",
    options=["📊 Parse & Download Excel", "🤖 Extract Text for AI"],
    horizontal=True,
    label_visibility="collapsed",
)
st.caption(
    "**Parse & Download Excel** — automatically extracts account details and all transactions into a styled Excel file."
    if mode == "📊 Parse & Download Excel"
    else "**Extract Text for AI** — pulls raw text from the PDF and builds a ready-to-paste prompt for ChatGPT, Claude, or Gemini."
)

st.divider()

# ── Bank Selector ─────────────────────────────────────────────────────────────
st.markdown("#### Select your Bank")
selected_bank = st.selectbox(
    label="bank",
    options=["UCO Bank", "Canara Bank", "SBI (Coming Soon)", "HDFC (Coming Soon)"],
    label_visibility="collapsed",
)

st.divider()

# ── Upload ────────────────────────────────────────────────────────────────────
uploaded_file = st.file_uploader(
    f"Upload your {selected_bank} statement PDF",
    type=["pdf"],
    help="Supports text-based PDFs. Ensure the selected bank matches the uploaded document."
)

pdf_password = st.text_input(
    "PDF password (optional)",
    type="password",
    help="Required only when the uploaded PDF is password protected.",
)

if not uploaded_file:
    st.info("👆 Upload a bank statement PDF to get started.")
    st.stop()

submit = st.button("🚀 Process PDF", type="primary", use_container_width=False)

if not submit:
    st.stop()

# BUG FIX: Read file bytes only ONCE
file_bytes = uploaded_file.read()

# ══════════════════════════════════════════════════════════════════════════════
# MODE A — Parse & Download Excel
# ══════════════════════════════════════════════════════════════════════════════
if mode == "📊 Parse & Download Excel":

    with st.spinner(f"Extracting data using {selected_bank} logic…"):
        try:
            page1_text = extract_first_page_text(file_bytes, pdf_password)
            
            # ROUTER CALLED HERE
            account_info, transactions = parse_bank_statement(selected_bank, file_bytes, page1_text, pdf_password)
            
            with open_pdf(file_bytes, pdf_password) as _pdf:
                page_count = len(_pdf.pages)
        except PdfPasswordError as e:
            st.error(str(e))
            st.stop()
        except CanaraOcrSetupError as e:
            st.error(str(e))
            st.stop()
        except Exception as e:
            st.error(f"Failed to read PDF: {e}")
            st.stop()

    if not transactions:
        if selected_bank == "Canara Bank":
            st.error(
                "No Canara transactions found. OCR ran, but the app could not detect the transaction rows. "
                "Check that the uploaded PDF is a Canara ePassbook/statement with the expected table layout, "
                "and that the scan is clear enough for OCR."
            )
        else:
            st.error("No transactions found. Make sure the PDF contains selectable text (not a scanned image).")
        st.stop()

    # ── Stats row ─────────────────────────────────────────────────────────────
    total_withdrawals = sum(safe_float(tx["withdrawal"]) for tx in transactions if tx.get("withdrawal"))
    total_deposits = sum(safe_float(tx["deposit"]) for tx in transactions if tx.get("deposit"))

    st.markdown(f"""
    <div class="stat-row">
      <div class="stat-card">
        <div class="stat-label">Transactions</div>
        <div class="stat-value">{len(transactions)}</div>
        <div class="stat-sub">{page_count} pages processed</div>
      </div>
      <div class="stat-card">
        <div class="stat-label">Total Withdrawals</div>
        <div class="stat-value" style="color:#C53030">₹{total_withdrawals:,.2f}</div>
        <div class="stat-sub">Money out</div>
      </div>
      <div class="stat-card">
        <div class="stat-label">Total Deposits</div>
        <div class="stat-value" style="color:#276749">₹{total_deposits:,.2f}</div>
        <div class="stat-sub">Money in</div>
      </div>
      <div class="stat-card">
        <div class="stat-label">Closing Balance</div>
        <div class="stat-value">₹{transactions[-1].get("balance","—")}</div>
        <div class="stat-sub">{transactions[-1].get("date","")}</div>
      </div>
    </div>
    """, unsafe_allow_html=True)

    # ── Download button ────────────────────────────────────────────────────────
    wb = build_excel_workbook(account_info, transactions, selected_bank)
    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)

    col_dl, col_spacer = st.columns([2, 5])
    with col_dl:
        st.download_button(
            label="⬇️ Download Excel",
            data=excel_buffer,
            file_name=f"{uploaded_file.name.replace('.pdf', '')}_statement.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
            type="primary",
        )

    st.divider()

    # ── Tabs ───────────────────────────────────────────────────────────────────
    tab_info, tab_tx = st.tabs(["🏛️ Account Details", f"💳 Transactions ({len(transactions)})"])

    with tab_info:
        if not any(account_info.values()):
            st.warning("Could not extract account details from this PDF automatically.")
        else:
            for label, value in account_info.items():
                if value:
                    col_label, col_value = st.columns([1, 2])
                    with col_label:
                        st.markdown(f"**{label}**")
                    with col_value:
                        st.code(value, language=None)

    with tab_tx:
        import pandas as pd

        df = pd.DataFrame([{
            "Date":            tx.get("date", ""),
            "Particulars":     tx.get("particulars", ""),
            "Withdrawal (₹)":  tx.get("withdrawal", ""),
            "Deposit (₹)":     tx.get("deposit", ""),
            "Balance (₹)":     tx.get("balance", ""),
        } for tx in transactions])

        st.markdown(f"**{len(transactions)} transactions** extracted from {page_count} pages")
        st.dataframe(
            df,
            use_container_width=True,
            height=520,
            hide_index=True,
            column_config={
                "Date":           st.column_config.TextColumn("Date", width=110),
                "Particulars":    st.column_config.TextColumn("Particulars", width=380),
                "Withdrawal (₹)": st.column_config.TextColumn("Withdrawal (₹)", width=130),
                "Deposit (₹)":    st.column_config.TextColumn("Deposit (₹)", width=130),
                "Balance (₹)":    st.column_config.TextColumn("Balance (₹)", width=140),
            }
        )

        empty_particulars = sum(1 for tx in transactions if not tx.get("particulars"))
        if empty_particulars > 0:
            st.markdown(f'<div class="warn-pill">⚠️ {empty_particulars} rows have empty Particulars — check column alignment</div>', unsafe_allow_html=True)
        else:
            st.markdown('<div class="success-pill">✅ All rows have Particulars</div>', unsafe_allow_html=True)

# ══════════════════════════════════════════════════════════════════════════════
# MODE B — Extract Text for AI
# ══════════════════════════════════════════════════════════════════════════════
else:
    with st.spinner("Extracting text from PDF…"):
        try:
            pdf_text, page_count = extract_pdf_text(file_bytes, pdf_password)
        except PdfPasswordError as e:
            st.error(str(e))
            st.stop()
        except Exception as e:
            st.error(f"Failed to read PDF: {e}")
            st.stop()

    if not pdf_text.strip():
        st.error("No extractable text found. This PDF may be scanned/image-based.")
        st.stop()

    prompt = AI_EXCEL_PROMPT_TEMPLATE.format(pdf_text=pdf_text)

    # ── Stats ──────────────────────────────────────────────────────────────────
    st.markdown(f"""
    <div class="stat-row">
      <div class="stat-card">
        <div class="stat-label">Pages</div>
        <div class="stat-value">{page_count}</div>
        <div class="stat-sub">extracted</div>
      </div>
      <div class="stat-card">
        <div class="stat-label">Characters</div>
        <div class="stat-value">{len(pdf_text):,}</div>
        <div class="stat-sub">raw text</div>
      </div>
      <div class="stat-card">
        <div class="stat-label">Prompt Size</div>
        <div class="stat-value">{len(prompt):,}</div>
        <div class="stat-sub">chars ready to paste</div>
      </div>
    </div>
    """, unsafe_allow_html=True)

    # ── Instructions ───────────────────────────────────────────────────────────
    st.info(
        "**How to use:** Download the prompt file below → open ChatGPT / Claude / Gemini "
        "→ paste the entire contents → the AI will return two CSV blocks (Account Details + Transactions) "
        "that you can copy into Excel."
    )

    # ── Download buttons ───────────────────────────────────────────────────────
    col1, col2, _spacer = st.columns([1, 1, 2])

    with col1:
        st.download_button(
            label="⬇️ Download Full Prompt (.txt)",
            data=prompt,
            file_name=f"{uploaded_file.name.replace('.pdf', '')}_ai_prompt.txt",
            mime="text/plain",
            use_container_width=True,
            type="primary",
        )
    with col2:
        st.download_button(
            label="⬇️ Download Raw Text Only",
            data=pdf_text,
            file_name=f"{uploaded_file.name.replace('.pdf', '')}_extracted_text.txt",
            mime="text/plain",
            use_container_width=True,
        )

    st.divider()

    # ── Preview tabs ───────────────────────────────────────────────────────────
    tab_prompt_prev, tab_text_prev = st.tabs(["📋 Prompt Preview", "📄 Raw Text Preview"])

    with tab_prompt_prev:
        st.caption("First 3000 characters of the full prompt")
        st.code(
            prompt[:3000] + ("\n\n… (truncated — download for full prompt)" if len(prompt) > 3000 else ""),
            language=None
        )

    with tab_text_prev:
        st.caption(f"First 3000 characters of extracted text ({page_count} pages total)")
        st.code(
            pdf_text[:3000] + ("\n\n… (truncated)" if len(pdf_text) > 3000 else ""),
            language=None
        )
